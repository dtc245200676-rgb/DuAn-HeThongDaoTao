import io, os
from datetime import timedelta
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from email_validator import validate_email, EmailNotValidError
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.api.deps import AuthContext, require_permission
from app.core.security import ACTIVATION_TOKEN_EXPIRE_HOURS, generate_one_time_token, generate_temporary_password, hash_one_time_token, hash_password, utcnow
from app.db.database import get_db
from app.models.role import Role
from app.models.tokens import ActivationToken
from app.models.user import User
from app.services.email_service import email_service
router=APIRouter(prefix="/api/users/import",tags=["S2-01 Import Excel"])
HEADERS=["full_name","email","phone","role_slugs"]
def _read_rows(raw:bytes):
    try: wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb.active
    except Exception as exc: raise HTTPException(422,"Không đọc được tệp Excel.") from exc
    try: first=[str(v or "").strip() for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
    except StopIteration: raise HTTPException(422,"Tệp Excel trống.")
    if first[:4]!=HEADERS: raise HTTPException(422,f"Dòng tiêu đề phải là: {', '.join(HEADERS)}")
    rows=[]
    for idx,values in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        row=dict(zip(HEADERS,[v if v is not None else "" for v in values[:4]]))
        if any(str(v).strip() for v in row.values()): rows.append((idx,row))
    return rows
def _valid_email(value:str)->bool:
    try: validate_email(value,check_deliverability=False); return True
    except EmailNotValidError: return False
def _validate_rows(db:Session,rows):
    role_map={r.slug:r for r in db.scalars(select(Role)).all()}; seen=set(); result=[]
    for excel_row,row in rows:
        errors=[]; name=str(row["full_name"]).strip(); email=str(row["email"]).strip().lower(); phone=str(row["phone"]).strip() or None
        roles=[x.strip() for x in str(row["role_slugs"] or "student").split(",") if x.strip()] or ["student"]
        if not name: errors.append("Thiếu họ tên")
        if not _valid_email(email): errors.append("Email không hợp lệ")
        if email in seen: errors.append("Email trùng trong tệp")
        seen.add(email)
        if email and db.scalar(select(User.id).where(User.email==email)): errors.append("Email đã tồn tại")
        invalid=[r for r in roles if r not in role_map]
        if invalid: errors.append("Vai trò không tồn tại: "+", ".join(invalid))
        result.append({"row":excel_row,"full_name":name,"email":email,"phone":phone,"role_slugs":roles,"errors":errors,"valid":not errors})
    return result,role_map
@router.get("/template")
def download_template(_:AuthContext=Depends(require_permission("users.import"))):
    wb=Workbook();ws=wb.active;ws.title="users";ws.append(HEADERS);ws.append(["Nguyễn Văn A","nguyenvana@example.com","0901234567","student"]);ws.append(["Trần Thị B","tranthib@example.com","0912345678","teacher"])
    stream=io.BytesIO();wb.save(stream);stream.seek(0)
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="mau-nhap-nguoi-dung.xlsx"'})
@router.post("/preview")
async def preview_import(file:UploadFile=File(...),_:AuthContext=Depends(require_permission("users.import")),db:Session=Depends(get_db)):
    if not (file.filename or "").lower().endswith(".xlsx"): raise HTTPException(422,"Chỉ chấp nhận tệp Excel .xlsx.")
    validated,_=_validate_rows(db,_read_rows(await file.read()))
    return {"rows":validated,"total":len(validated),"valid":sum(r["valid"] for r in validated),"invalid":sum(not r["valid"] for r in validated)}
@router.post("")
async def import_users(file:UploadFile=File(...),_:AuthContext=Depends(require_permission("users.import")),db:Session=Depends(get_db)):
    if not (file.filename or "").lower().endswith(".xlsx"): raise HTTPException(422,"Chỉ chấp nhận tệp Excel .xlsx.")
    validated,role_map=_validate_rows(db,_read_rows(await file.read()));imported=[];skipped=[];now=utcnow()
    for row in validated:
        if not row["valid"]: skipped.append(row);continue
        temp=generate_temporary_password(); raw_token=generate_one_time_token(); roles=[role_map[r] for r in row["role_slugs"]]
        user=User(full_name=row["full_name"],email=row["email"],phone=row["phone"],password_hash=hash_password(temp),role=roles[0].slug,status="pending",is_active=False,must_change_password=True,roles=roles)
        db.add(user);db.flush();db.add(ActivationToken(user_id=user.id,token_hash=hash_one_time_token(raw_token),expires_at=now+timedelta(hours=ACTIVATION_TOKEN_EXPIRE_HOURS)))
        activation_url=f"{os.getenv('FRONTEND_URL','http://localhost:5173')}/activate?token={raw_token}"
        sent=email_service.send(user.email,"Kích hoạt tài khoản - Hệ thống đào tạo CodeGym",f"Tài khoản của bạn đã được tạo từ danh sách Excel.\nMật khẩu tạm: {temp}\nLiên kết kích hoạt: {activation_url}")
        imported.append({"row":row["row"],"id":user.id,"email":user.email,"email_sent":sent})
    db.commit();return {"message":"Đã hoàn tất nhập danh sách.","total":len(validated),"imported_count":len(imported),"skipped_count":len(skipped),"imported":imported,"skipped":skipped}
